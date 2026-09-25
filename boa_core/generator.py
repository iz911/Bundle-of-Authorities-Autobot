"""
BOA Document Builder.
Constructs Microsoft Word (.docx) Bundle of Authorities
matching Singapore Supreme Court and State Courts Practice Directions.
"""
import re
from typing import Dict, List, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from .citations import own_paragraph_number, paragraph_numbers
from .models import AuthorityItem, CaseMetadata, RetrievalStatus


def set_cell_border(cell, **kwargs):
    """
    Sets cell borders.
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='000000')
    """
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'<w:top w:val="{kwargs.get("top", {}).get("val", "single")}" w:sz="{kwargs.get("top", {}).get("sz", "4")}" w:space="0" w:color="{kwargs.get("top", {}).get("color", "000000")}"/>\n'
        f'<w:left w:val="{kwargs.get("left", {}).get("val", "single")}" w:sz="{kwargs.get("left", {}).get("sz", "4")}" w:space="0" w:color="{kwargs.get("left", {}).get("color", "000000")}"/>\n'
        f'<w:bottom w:val="{kwargs.get("bottom", {}).get("val", "single")}" w:sz="{kwargs.get("bottom", {}).get("sz", "4")}" w:space="0" w:color="{kwargs.get("bottom", {}).get("color", "000000")}"/>\n'
        f'<w:right w:val="{kwargs.get("right", {}).get("val", "single")}" w:sz="{kwargs.get("right", {}).get("sz", "4")}" w:space="0" w:color="{kwargs.get("right", {}).get("color", "000000")}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)


def set_cell_background(cell, hex_color: str):
    """Sets background shading of a cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def build_bundle_of_authorities(
    metadata: CaseMetadata,
    grouped_authorities: Dict[str, List[AuthorityItem]],
    output_docx_path: str
) -> str:
    """
    Builds the complete Bundle of Authorities .docx file.
    """
    doc = docx.Document()

    # 1. Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 2. Cover Page
    _create_cover_page(doc, metadata)

    # 3. Index to Bundle of Authorities
    doc.add_page_break()
    _create_index_page(doc, metadata, grouped_authorities)

    # 4. Tabbed Authorities & Placeholders
    tab_number = 1
    for group_name, items in grouped_authorities.items():
        for auth in items:
            doc.add_page_break()
            _create_tab_section(doc, tab_number, group_name, auth)
            tab_number += 1

    # Save output docx
    doc.save(output_docx_path)
    return output_docx_path


def _create_cover_page(doc: docx.Document, meta: CaseMetadata):
    """Generates the cover page matching Singapore court practice."""
    # Court Name
    p_court = doc.add_paragraph()
    p_court.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_court = p_court.add_run(meta.court_name.upper())
    r_court.bold = True
    r_court.font.name = "Times New Roman"
    r_court.font.size = Pt(12)

    # Case / Summons Number
    if meta.case_number:
        p_num = doc.add_paragraph()
        p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for line in meta.case_number.strip().split("\n"):
            r_num = p_num.add_run(line.strip() + "\n")
            r_num.bold = True
            r_num.font.name = "Times New Roman"
            r_num.font.size = Pt(12)

    # In the matter of...
    if meta.matter_description:
        p_matter = doc.add_paragraph()
        p_matter.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_matter = p_matter.add_run(meta.matter_description)
        r_matter.italic = True
        r_matter.font.name = "Times New Roman"
        r_matter.font.size = Pt(11)

    # Spacing
    doc.add_paragraph()

    # Parties Section
    p_between = doc.add_paragraph()
    p_between.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_b = p_between.add_run("Between\n\n")
    r_b.font.name = "Times New Roman"

    # Claimant
    r_c_name = p_between.add_run(f"{meta.claimant_name.upper()}\n")
    r_c_name.bold = True
    r_c_name.font.name = "Times New Roman"
    if meta.claimant_uen:
        r_c_uen = p_between.add_run(f"(Singapore UEN No. {meta.claimant_uen})\n")
        r_c_uen.font.name = "Times New Roman"
    r_c_role = p_between.add_run("...Claimant\n\nAnd\n\n")
    r_c_role.font.name = "Times New Roman"

    # Respondent
    r_r_name = p_between.add_run(f"{meta.respondent_name.upper()}\n")
    r_r_name.bold = True
    r_r_name.font.name = "Times New Roman"
    if meta.respondent_uen:
        r_r_uen = p_between.add_run(f"(Singapore UEN No. {meta.respondent_uen})\n")
        r_r_uen.font.name = "Times New Roman"
    r_r_role = p_between.add_run("...Respondent\n\n")
    r_r_role.font.name = "Times New Roman"

    # Spacing
    doc.add_paragraph()

    # Document Title (e.g. CLAIMANT'S BUNDLE OF AUTHORITIES)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(meta.document_title.upper())
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)

    # Date
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_date = p_date.add_run(meta.dated_date)
    r_date.font.name = "Times New Roman"
    r_date.font.size = Pt(11)

    # Spacing
    doc.add_paragraph()

    # 2-Column Solicitors Table
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell_left = table.cell(0, 0)
    cell_right = table.cell(0, 1)
    cell_left.width = Inches(3.2)
    cell_right.width = Inches(3.2)

    # Fill Solicitors Info
    p_sol_c = cell_left.paragraphs[0]
    p_sol_c.add_run(meta.claimant_solicitors.strip()).font.size = Pt(9.5)

    p_sol_r = cell_right.paragraphs[0]
    p_sol_r.add_run(meta.respondent_solicitors.strip()).font.size = Pt(9.5)


def _create_index_page(
    doc: docx.Document,
    meta: CaseMetadata,
    grouped_authorities: Dict[str, List[AuthorityItem]]
):
    """Generates the Index table matching Singapore BOA conventions."""
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_type = meta.document_title.replace("BUNDLE OF AUTHORITIES", "").strip()
    index_title = f"INDEX TO {meta.document_title.upper()}"
    r_head = p_head.add_run(index_title)
    r_head.bold = True
    r_head.font.name = "Times New Roman"
    r_head.font.size = Pt(13)

    doc.add_paragraph()

    # Table with TAB and DESCRIPTION
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].width = Inches(1.0)
    hdr_cells[1].width = Inches(5.5)

    r_tab = hdr_cells[0].paragraphs[0].add_run("TAB")
    r_tab.bold = True
    r_tab.font.name = "Times New Roman"

    r_desc = hdr_cells[1].paragraphs[0].add_run("DESCRIPTION")
    r_desc.bold = True
    r_desc.font.name = "Times New Roman"

    set_cell_background(hdr_cells[0], "F2F2F2")
    set_cell_background(hdr_cells[1], "F2F2F2")
    set_cell_border(hdr_cells[0])
    set_cell_border(hdr_cells[1])

    tab_counter = 1

    for group_name, items in grouped_authorities.items():
        if not items:
            continue

        # Category Section Heading Row
        row_grp = table.add_row()
        c_grp_tab = row_grp.cells[0]
        c_grp_desc = row_grp.cells[1]
        c_grp_tab.width = Inches(1.0)
        c_grp_desc.width = Inches(5.5)

        r_grp_title = c_grp_desc.paragraphs[0].add_run(group_name)
        r_grp_title.bold = True
        r_grp_title.italic = True
        r_grp_title.font.name = "Times New Roman"
        r_grp_title.font.size = Pt(11)

        set_cell_background(c_grp_tab, "EAEAEA")
        set_cell_background(c_grp_desc, "EAEAEA")
        set_cell_border(c_grp_tab)
        set_cell_border(c_grp_desc)

        # Authority Rows
        for auth in items:
            row = table.add_row()
            c_tab = row.cells[0]
            c_desc = row.cells[1]
            c_tab.width = Inches(1.0)
            c_desc.width = Inches(5.5)

            # Tab number
            p_tab = c_tab.paragraphs[0]
            p_tab.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_tnum = p_tab.add_run(str(tab_counter))
            r_tnum.bold = True
            r_tnum.font.name = "Times New Roman"

            # Description (Title, Citation, Pinpoints, Relevance)
            p_desc = c_desc.paragraphs[0]
            
            # Title & Citation
            citation_display = f"{auth.title} {auth.citation or ''}".strip()
            r_cit = p_desc.add_run(citation_display + "\n")
            r_cit.bold = True
            r_cit.font.name = "Times New Roman"
            r_cit.font.size = Pt(10.5)

            # Relevance / Pinpoints
            if auth.relevance:
                r_rel = p_desc.add_run(f"Relevance: {auth.relevance}\n")
                r_rel.italic = True
                r_rel.font.name = "Times New Roman"
                r_rel.font.size = Pt(9.5)
            elif auth.pinpoints:
                pin_str = ", ".join(auth.pinpoints)
                r_pin = p_desc.add_run(f"Cited Pinpoint(s): {pin_str}\n")
                r_pin.font.name = "Times New Roman"
                r_pin.font.size = Pt(9.5)

            # Status Note: every authority without verified text gets a placeholder tab
            if auth.retrieval_status != RetrievalStatus.VERIFIED_MATCH or not auth.retrieved_text:
                r_status = p_desc.add_run("[Manual Tab Insertion Required]")
                r_status.font.size = Pt(8.5)
                r_status.font.color.rgb = RGBColor(180, 50, 50)

            set_cell_border(c_tab)
            set_cell_border(c_desc)

            tab_counter += 1


def _create_tab_section(
    doc: docx.Document,
    tab_num: int,
    group_name: str,
    auth: AuthorityItem
):
    """Creates a distinct tab section in the BOA."""
    # Tab Header: e.g. "Tab-1" or "TAB 1"
    p_tab_hdr = doc.add_paragraph()
    r_tab_hdr = p_tab_hdr.add_run(f"Tab-{tab_num}")
    r_tab_hdr.bold = True
    r_tab_hdr.font.name = "Times New Roman"
    r_tab_hdr.font.size = Pt(16)

    # Authority Heading
    p_auth_hdr = doc.add_paragraph()
    r_auth_title = p_auth_hdr.add_run(f"{auth.title}\n")
    r_auth_title.bold = True
    r_auth_title.font.name = "Times New Roman"
    r_auth_title.font.size = Pt(13)

    if auth.citation:
        r_auth_cit = p_auth_hdr.add_run(f"{auth.citation}\n")
        r_auth_cit.bold = True
        r_auth_cit.font.name = "Times New Roman"
        r_auth_cit.font.size = Pt(11)

    # Pinpoints
    if auth.pinpoints:
        p_pin = doc.add_paragraph()
        r_pin_lbl = p_pin.add_run("Cited Pinpoint(s): ")
        r_pin_lbl.bold = True
        r_pin_lbl.font.size = Pt(10)
        r_pin_val = p_pin.add_run(", ".join(auth.pinpoints))
        r_pin_val.font.size = Pt(10)
        # Highlight pinpoints banner in yellow
        r_pin_val.font.highlight_color = WD_COLOR_INDEX.YELLOW

    doc.add_paragraph()

    # Body Content: Verified vs Placeholder (IDK)
    if auth.retrieval_status == RetrievalStatus.VERIFIED_MATCH and auth.retrieved_text:
        # Split text into paragraphs and highlight exactly the judgment's own cited ones, never a quotation
        cited = paragraph_numbers(auth.pinpoints)
        paragraphs = auth.retrieved_text.split("\n")
        for para in paragraphs:
            para_clean = para.strip()
            if not para_clean:
                continue

            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)

            should_highlight = own_paragraph_number(para_clean) in cited

            run = p.add_run(para_clean)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            if should_highlight:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    else:
        # IDK Placeholder Page
        _create_placeholder_page(doc, tab_num, auth)


def _create_placeholder_page(doc: docx.Document, tab_num: int, auth: AuthorityItem):
    """Generates a standardized, clean placeholder tab page for counsel to insert."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)

    set_cell_background(cell, "FFF9E6") # Soft warning yellow
    set_cell_border(cell, sz="8", color="CC9900")

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)

    r_warn = p.add_run("MANUAL INSERTION REQUIRED\n\n")
    r_warn.bold = True
    r_warn.font.name = "Times New Roman"
    r_warn.font.size = Pt(12)
    r_warn.font.color.rgb = RGBColor(160, 80, 0)

    r_desc = p.add_run(
        f"Tab Number: Tab {tab_num}\n"
        f"Authority: {auth.title}\n"
        f"Citation: {auth.citation or auth.raw_citation}\n"
        f"Category: {auth.category.value}\n"
        f"Cited Pinpoint(s): {', '.join(auth.pinpoints) if auth.pinpoints else 'Whole authority cited'}\n\n"
    )
    r_desc.font.name = "Times New Roman"
    r_desc.font.size = Pt(10.5)

    r_reason_lbl = p.add_run("Why this is a placeholder:\n")
    r_reason_lbl.bold = True
    r_reason_lbl.font.size = Pt(10)
    r_reason = p.add_run(f"{auth.retrieval_note or 'It was not checked against an official source. Insert it by hand.'}\n\n")
    r_reason.font.size = Pt(10)

    r_instr_lbl = p.add_run("Instructions for Litigation Counsel / Legal Trainee:\n")
    r_instr_lbl.bold = True
    r_instr_lbl.font.size = Pt(10)

    instructions = (
        "1. Retrieve the official report/statute from LawNet, Westlaw, ICLR, or firm textbook.\n"
        "2. Print or insert the scanned pages directly behind this Tab divider.\n"
        "3. Highlight the cited pinpoint passages in yellow per Supreme Court Practice Directions."
    )
    r_instr = p.add_run(instructions)
    r_instr.italic = True
    r_instr.font.size = Pt(9.5)
